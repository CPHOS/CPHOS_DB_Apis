# DO NOT USE THIS SCRIPT FOR INDIVIDUAL ENROLLMENT!!!!!!

'''
Excel表格格式要求：
第一行为表头，从第二行开始处理。
第 1 列：姓名。
第 3 列：微信小程序昵称 / 虚拟 user id。
第 5 列：学校。
第 6 列：省份或地区，脚本中取前两位作为地区名。
本脚本用于审核领队教练，teacher_sup 固定为 0，teacher_students 固定为 100。
上传版本不填写真实 xlsx 路径，运行前在 INPUT_XLSX_PATH 和 OUTPUT_XLSX_PATH 中填入本地路径。
'''

from db_api import customTransaction
from db_api.DataQueryApis.GetTeacherInfoApis import *
from db_api.DataManagingApis.ChangeTeacherInfoApis import *
from db_api.DataManagingApis.ChangeSchoolInfoApis import *
from db_api.DataQueryApis.GetSchoolInfoApis import *
from db_api.DataQueryApis.GetStudentInfoApis import *
from db_api.DataQueryApis.GetAreaInfoApis import *

from openpyxl import load_workbook, Workbook

INPUT_XLSX_PATH = ''
OUTPUT_XLSX_PATH = ''

wb = load_workbook(INPUT_XLSX_PATH)
ws = wb.active
wbb = Workbook()
wss = wbb.active
wss.cell(1, 1).value = '微信名'
wss.cell(1, 2).value = '姓名'
wss.cell(1, 3).value = '结果'
wss.cell(1, 4).value = '注释'
# wss.cell(1, 6).value = '上级'

rowmax = ws.max_row
for i in range(2, rowmax + 1):
    user_nickname = ws.cell(row=i, column=3).value
    teacher_name = ws.cell(row=i, column=1).value
    teacher_school = ws.cell(row=i, column=5).value
    teacher_sup = 0
    teacher_area = ws.cell(row=i, column=6).value
    teacher_area = teacher_area[:2]
    teacher_students = 100
    print("正在处理名为{}的{}，学校为{}，地区为{}".format(user_nickname, teacher_name, teacher_school, teacher_area))

    wss.cell(i, 1).value = user_nickname
    wss.cell(i, 2).value = teacher_name

    already_exist_flag = False
    already_exist_teacher = customTransaction.executeOperation(GetTeacherInfoByName(teacher_name))
    if (len(already_exist_teacher) > 1):
        wss.cell(i, 3).value = '已存在多个，似乎有点问题'
        wss.cell(i, 4).value = '已存在多个，似乎有点问题'
        continue
    elif (len(already_exist_teacher) == 1):
        print("已经存在！".format(teacher_name))
        already_exist_flag = True
        teacher_id = already_exist_teacher[0]['id']

    if already_exist_flag:
        wss.cell(i, 7).value = '已存在'
    else:
        wss.cell(i, 7).value = '本次审核前不存在，审核后是否存在请参看前几列。'

    if already_exist_flag:
        teacher_p_id = already_exist_teacher[0]['p_id']
        if teacher_p_id != 0:
            customTransaction.executeOperation(MakeAllTypesToBeSupTeacher(teacher_id))
            wss.cell(i, 7).value = '已存在，但是之前不是领队，现在改成领队了。'
            print("叫做{}的老师之前不是领队，现在已经改成领队了".format(teacher_name))

    if not already_exist_flag:
        teacher_id_list = customTransaction.executeOperation(GetToBeVerifiedTeacherInfoByWechatName(user_nickname))
        if (len(teacher_id_list) == 0):
            print("故障：未找到微信名为{}，名称为{}的待审核教练".format(user_nickname, teacher_name))
            wss.cell(i, 3).value = '未找到'
            wss.cell(i, 4).value = '未提交审核'
            continue
        elif (len(teacher_id_list) > 1):
            wss.cell(i, 3).value = '故障'
            wss.cell(i, 4).value = '有重复的微信名或提交多份审核'
            continue
        teacher_id = teacher_id_list[0]['id']

        school_id_list = customTransaction.executeOperation(GetSchoolInfoByName(teacher_school))
        if (len(school_id_list) == 0):
            area_id_lst = customTransaction.executeOperation(GetAreaInfoByName(teacher_area))
            if (len(area_id_lst) == 1):
                area_id = area_id_lst[0]['id']
            elif (len(area_id_lst) > 1):
                wss.cell(i, 3).value = '故障'
                wss.cell(i, 4).value = '数据库有重复地区名，理应不会发生'
                continue
            else:
                wss.cell(i, 3).value = '故障'
                wss.cell(i, 4).value = '地区名输入错误'
                continue
            school_id = customTransaction.executeOperation(AddNewSchoolByName(teacher_school, area_id))
            wss.cell(i, 5).value = '我们新增了一个叫做{}的，在{}学校'.format(teacher_school, teacher_area)
            print('我们新增了一个叫做{}的，在{}学校'.format(teacher_school, teacher_area))
        elif (len(school_id_list) > 1):
            wss.cell(i, 3).value = '故障'
            wss.cell(i, 4).value = '有重名的学校'
            continue
        else:
            school_id = school_id_list[0]['id']

        flag = 0
        verified_teacher_lst = customTransaction.executeOperation(GetTeacherInfoByName(teacher_name))
        for item in verified_teacher_lst:
            if item['wechat_nickname'] == user_nickname:
                flag = 1
                break
        if flag:
            wss.cell(i, 3).value = '故障'
            wss.cell(i, 4).value = '已完成审核'
            continue

        customTransaction.executeOperation(
            VerifyTeacherUserToBeSupTeacher(teacher_id, teacher_name, school_id, 1, teacher_students))

    if (type(teacher_sup) != type(0)):
        verified_teacher_id = teacher_id
        supteacher_lst = customTransaction.executeOperation(
            GetTeacherInfoByName(teacher_sup))
        if len(supteacher_lst) == 0:
            wss.cell(i, 3).value = '故障'
            wss.cell(i, 4).value = '审核已通过，但负责人不存在'
            continue
        elif len(supteacher_lst) > 1:
            wss.cell(i, 3).value = '故障'
            wss.cell(i, 4).value = '审核已通过，但有重名的上级负责人'
            continue
        supteacher_id = supteacher_lst[0]['id']
        customTransaction.executeOperation(
            MakeAllTypesToBeSubCoach(ChangedUserId=verified_teacher_id, ItsNewSupId=supteacher_id))
    wss.cell(i, 3).value = '通过！'

wbb.save(OUTPUT_XLSX_PATH)

# CAUTION WHEN COMMIT!!!!
# customTransaction.commit()
