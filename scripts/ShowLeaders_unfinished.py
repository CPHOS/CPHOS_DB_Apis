from db_api import customTransaction
from db_api.DataQueryApis.GetTeacherInfoApis import *
from db_api.DataManagingApis.ChangeTeacherInfoApis import *
from db_api.DataManagingApis.ChangeSchoolInfoApis import *
from db_api.DataQueryApis.GetSchoolInfoApis import *
from db_api.DataQueryApis.GetStudentInfoApis import *
from db_api.DataQueryApis.GetAreaInfoApis import *

from openpyxl import load_workbook, Workbook
from tqdm import tqdm

wbb = Workbook()
wss = wbb.active
wss.cell(1, 1).value = '题目'
wss.cell(1, 2).value = '姓名'
wss.cell(1, 3).value = '待批阅数量'

all_teacher = customTransaction.executeOperation(GetTeacherInfoByFlexibleName(''))
i = 2

for teacher in tqdm(all_teacher):
    teacher_id = teacher['id']
    NotViewed_number = customTransaction.executeOperation(GetTeacherNotViewdProblemNumber(teacher_id))
    number =  NotViewed_number[0]

    if number != 0 :
        wss.cell(i,3).value = number
        wss.cell(i,1).value = teacher['viewing_problem']
        wss.cell(i,2).value = teacher['user_name']
        i = i+1

wbb.save('')
