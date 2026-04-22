# DO NOT USE THIS SCRIPT FOR TEAM ENROLLMENT!!!!!!

'''
Excel表格格式要求：
第一行为表头，从第二行开始处理。
第 1 列：姓名。
第 3 列：微信小程序昵称 / 虚拟 user id。
第 5 列：学校。本脚本沿用原逻辑，实际统一使用数据库中的“个人”学校。
第 6 列：地区名，当数据库中不存在“个人”学校时用于新增学校。
本脚本用于个人报名审核，teacher_students 固定为 1。
上传版本不填写真实 xlsx 路径，运行前在 INPUT_XLSX_PATH 和 OUTPUT_XLSX_PATH 中填入本地路径。
'''

from db_api import customTransaction
from db_api.DataQueryApis.GetTeacherInfoApis import *
from db_api.DataManagingApis.ChangeTeacherInfoApis import *
from db_api.DataManagingApis.ChangeSchoolInfoApis import *
from db_api.DataQueryApis.GetSchoolInfoApis import *
from db_api.DataQueryApis.GetStudentInfoApis import *
from db_api.DataQueryApis.GetAreaInfoApis import *

from openpyxl import load_workbook, Workbook

INPUT_XLSX_PATH = ''
OUTPUT_XLSX_PATH = ''

wb = load_workbook(INPUT_XLSX_PATH)
ws = wb.active
wbb = Workbook()
wss = wbb.active
wss.cell(1, 1).value = '微信名'
wss.cell(1, 2).value = '姓名'
wss.cell(1, 3).value = '结果'
wss.cell(1, 4).value = '注释'

rowmax = ws.max_row
for i in range(2, rowmax + 1):
    user_nickname = ws.cell(row=i, column=3).value  # user_nickname小程序中使用的是虚拟id，注意使用虚拟id
    teacher_name = ws.cell(row=i, column=1).value
    # teacher_school = ws.cell(row=i, column=5).value
    teacher_area = ws.cell(row=i, column=6).value
    teacher_area = teacher_area
    teacher_students = 1

    wss.cell(i, 1).value = user_nickname
    wss.cell(i, 2).value = teacher_name

    flag = 0
    verified_teacher_lst = customTransaction.executeOperation(GetTeacherInfoByName(teacher_name))
    for item in verified_teacher_lst:
        print(item['user_name'], teacher_name)
        if item['wechat_nickname'] == user_nickname:
            flag = 1
            break
    if flag:
        wss.cell(i, 3).value = '存在'
        wss.cell(i, 4).value = '已完成审核'
        continue

    teacher_id_list = customTransaction.executeOperation(GetToBeVerifiedTeacherInfoByWechatName(user_nickname))
    if (len(teacher_id_list) == 0):
        wss.cell(i, 3).value = '故障'
        wss.cell(i, 4).value = '未提交审核'
        continue
    elif (len(teacher_id_list) > 1):
        wss.cell(i, 3).value = '故障'
        wss.cell(i, 4).value = '有重复的微信名或提交多份审核'
        continue
    teacher_id = teacher_id_list[0]['id']

    school_id_list = customTransaction.executeOperation(GetSchoolInfoByName('个人'))
    if (len(school_id_list) == 0):
        area_id_lst = customTransaction.executeOperation(GetAreaInfoByName(teacher_area))
        if (len(area_id_lst) == 1):
            area_id = area_id_lst[0]['id']
        elif (len(area_id_lst) > 1):
            wss.cell(i, 3).value = '故障'
            wss.cell(i, 4).value = '数据库有重复地区名，理应不会发生'
            continue
        else:
            wss.cell(i, 3).value = '故障'
            wss.cell(i, 4).value = '地区名输入错误'
            continue
        school_id = customTransaction.executeOperation(AddNewSchoolByName('个人', area_id))
    elif (len(school_id_list) > 1):
        wss.cell(i, 3).value = '故障'
        wss.cell(i, 4).value = '有重名的学校'
        continue
    else:
        school_id = school_id_list[0]['id']

    wss.cell(i, 3).value = '通过'
    customTransaction.executeOperation(VerifyTeacherUserToBeSupTeacher(teacher_id, teacher_name, school_id, 1, 1))

wbb.save(OUTPUT_XLSX_PATH)

# CAUTION WHEN COMMIT!!!!
# customTransaction.commit()
